"""Bounded, credential-free exports and protected short-lived downloads."""
from __future__ import annotations
from pathlib import Path
import csv
import io
import re
import secrets
import os
import time
from webapi.runtime37 import connection, utcnow, sqlite_snapshot

TOKEN_TTL_SECONDS = 3600
FILE_RETENTION_SECONDS = 7 * 86400
GENERATED_PREFIXES = (
    'web_report_', 'web_db_', 'NetworkAutomation_DR_PRIVATE_',
    'diagnostics_', 'device_'
)


def _relative_display(path: Path) -> str:
    from app_runtime import DATA_DIR
    try:
        return str(path.resolve().relative_to(Path(DATA_DIR).resolve())).replace('\\', '/')
    except Exception:
        return path.name




def _export_root(destination_token: str | None = None) -> Path:
    if destination_token:
        from desktop_export import consume_export_directory
        chosen = consume_export_directory(destination_token)
        if chosen is None:
            raise ValueError('EXPORT_DIRECTORY_SELECTION_REQUIRED')
        return chosen
    from app_runtime import REPORT_DIR
    return Path(REPORT_DIR)


def _safe_export_name(filename: str, extension: str) -> str:
    name = Path(str(filename or ('NetworkAutomation' + extension))).name
    name = re.sub(r'[^0-9A-Za-z._-]+', '_', name).strip('._') or ('NetworkAutomation' + extension)
    if not name.lower().endswith(extension.lower()):
        name = Path(name).stem + extension
    return name


def _unique_output(root: Path, filename: str) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    candidate = root / filename
    if not candidate.exists():
        return candidate
    stamp = time.strftime('%Y%m%d_%H%M%S')
    base, suffix = candidate.stem, candidate.suffix
    candidate = root / f'{base}_{stamp}{suffix}'
    idx = 2
    while candidate.exists():
        candidate = root / f'{base}_{stamp}_{idx}{suffix}'
        idx += 1
    return candidate


def cleanup_generated_files(now=None):
    from app_runtime import REPORT_DIR, BACKUP_DIR
    cutoff = float(now or time.time()) - FILE_RETENTION_SECONDS
    for root in (Path(REPORT_DIR), Path(BACKUP_DIR)):
        if not root.exists():
            continue
        for p in root.iterdir():
            try:
                if p.is_file() and p.name.startswith(GENERATED_PREFIXES) and p.stat().st_mtime < cutoff:
                    p.unlink(missing_ok=True)
            except OSError:
                pass


def ensure_registry():
    with connection() as c:
        c.execute('CREATE TABLE IF NOT EXISTS web_downloads37(token TEXT PRIMARY KEY,name TEXT,path TEXT,owner_id INTEGER,kind TEXT,created_at TEXT,expires_epoch REAL)')
        cols = {x['name'] for x in c.execute('PRAGMA table_info(web_downloads37)')}
        if 'expires_epoch' not in cols:
            c.execute('ALTER TABLE web_downloads37 ADD COLUMN expires_epoch REAL')
        c.execute('DELETE FROM web_downloads37 WHERE expires_epoch IS NOT NULL AND expires_epoch<?', (time.time(),))
    cleanup_generated_files()


def register(path, owner_id, kind):
    ensure_registry()
    path = Path(path)
    token = secrets.token_urlsafe(24)
    expires = time.time() + TOKEN_TTL_SECONDS
    with connection() as c:
        c.execute(
            'INSERT INTO web_downloads37(token,name,path,owner_id,kind,created_at,expires_epoch) VALUES(?,?,?,?,?,?,?)',
            (token, path.name, str(path.resolve()), owner_id, kind, utcnow(), expires),
        )
    return {
        'success': True,
        'name': path.name,
        'download_url': '/api/downloads/' + token,
        'saved_relative_path': _relative_display(path),
        'expires_in_seconds': TOKEN_TTL_SECONDS,
    }


def backup_db(owner_id):
    from database.db import DB_PATH
    from app_runtime import BACKUP_DIR
    path = Path(BACKUP_DIR) / ('web_db_' + secrets.token_hex(12) + '.db')
    sqlite_snapshot(Path(DB_PATH), path)
    return register(path, owner_id, 'database')


def export_report(owner_id, destination_token: str | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from webapi.data37 import merged_devices

    wb = Workbook()
    meta = wb.active
    meta.title = 'Readme'
    meta.append(['NetworkAutomation operational export', utcnow()])
    meta.append(['Source', 'Selected desktop database / local API'])
    meta.append(['Scope', 'Maximum 5000 latest rows per sheet; no credentials/configuration content'])
    meta.append(['Meaning', 'Unknown/Stale is not Offline; status is time-bound'])
    tables = {'Devices': merged_devices()[:5000]}
    with connection() as c:
        for tab, table, fields in [
            ('Alerts', 'alerts', ['id', 'ip', 'ip_address', 'severity', 'message', 'status', 'resolved', 'created_at']),
            ('Server targets', 'server_monitor_targets', ['id', 'name', 'host', 'port', 'protocol', 'enabled']),
            ('Ping history', 'ping_results', ['ip', 'ip_address', 'status', 'response', 'response_ms', 'ping_time']),
            ('Health history', 'health_samples', ['host', 'cpu', 'memory', 'latency_ms', 'packet_loss', 'created_at']),
        ]:
            cols = {x['name'] for x in c.execute(f'PRAGMA table_info("{table}")')}
            use = [f for f in fields if f in cols]
            if use:
                tables[tab] = [dict(x) for x in c.execute('SELECT ' + ','.join(use) + f' FROM "{table}" ORDER BY rowid DESC LIMIT 5000')]
    for name, rows in tables.items():
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(['No data'])
            continue
        cols = list(rows[0])
        ws.append(cols)
        for row in rows:
            values = []
            for key in cols:
                v = row.get(key)
                if isinstance(v, str):
                    v = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', v)[:32000]
                    if v.lstrip().startswith(('=', '+', '-', '@')):
                        v = "'" + v
                values.append(v)
            ws.append(values)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    for ws in wb:
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='183B56')
        ws.row_dimensions[1].height = 26
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(color='176B45')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    root = _export_root(destination_token)
    out = _unique_output(root, 'NetworkAutomation_Report_' + time.strftime('%Y%m%d_%H%M%S') + '.xlsx')
    wb.save(out)
    return register(out, owner_id, 'report')


def save_csv(owner_id, filename: str, content: str, destination_token: str | None = None):
    if not isinstance(content, str):
        raise ValueError('CSV_CONTENT_INVALID')
    if len(content.encode('utf-8')) > 2 * 1024 * 1024:
        raise ValueError('CSV_TOO_LARGE')
    stem = _safe_export_name(filename, '.csv')
    out = _unique_output(_export_root(destination_token), stem)
    text = content if content.startswith('\ufeff') else '\ufeff' + content
    tmp = out.with_suffix(out.suffix + '.tmp')
    tmp.write_text(text, encoding='utf-8', newline='')
    os.replace(tmp, out)
    return register(out, owner_id, 'csv')


def save_xlsx_from_csv(owner_id, filename: str, content: str, destination_token: str | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if not isinstance(content, str):
        raise ValueError('XLSX_CONTENT_INVALID')
    if len(content.encode('utf-8')) > 2 * 1024 * 1024:
        raise ValueError('XLSX_SOURCE_TOO_LARGE')
    rows = list(csv.reader(io.StringIO(content.lstrip('\ufeff'))))
    if len(rows) > 10000:
        raise ValueError('XLSX_TOO_MANY_ROWS')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dữ liệu'
    for row in rows:
        safe_row = []
        for value in row:
            value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))[:32000]
            if value.lstrip().startswith(('=', '+', '-', '@')):
                value = "'" + value
            safe_row.append(value)
        ws.append(safe_row)
    if rows:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='183B56')
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    out = _unique_output(_export_root(destination_token), _safe_export_name(filename, '.xlsx'))
    tmp = out.with_suffix(out.suffix + '.tmp')
    wb.save(tmp)
    os.replace(tmp, out)
    return register(out, owner_id, 'xlsx')
