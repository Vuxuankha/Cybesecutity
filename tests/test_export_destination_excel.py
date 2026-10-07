from pathlib import Path
import tempfile


def text(name):
    return Path(name).read_text(encoding='utf-8')


def test_native_folder_picker_bridge_is_wired():
    launcher = text('desktop_launcher.py')
    assert 'class _DesktopJsApi' in launcher
    assert 'choose_export_directory' in launcher
    assert 'webview.FOLDER_DIALOG' in launcher
    assert 'js_api=js_api' in launcher
    assert 'js_api.window = window' in launcher


def test_ui_offers_csv_and_excel_with_destination_picker():
    app = text('webapi/static/app.js')
    wb = text('webapi/static/workbench45.js')
    cyber = text('webapi/static/cybersecurity51.js')
    assert 'chooseExportDirectory' in app
    assert "button('Xu\\u1ea5t CSV...'" in wb
    assert "button('Xu\\u1ea5t Excel...'" in wb
    assert "'/reports/xlsx':'/reports/csv'" in wb
    assert 'destination_token:picked.token' in wb
    assert "button('Excel...'" in cyber
    assert "'/reports/xlsx':'/reports/csv'" in cyber


def test_export_destination_token_is_one_time_and_files_land_in_selected_folder(monkeypatch):
    import desktop_export
    import webapi.reports37 as reports
    from openpyxl import load_workbook

    monkeypatch.setattr(reports, 'register', lambda path, owner_id, kind: {'success': True, 'name': Path(path).name, 'path': str(path), 'kind': kind})
    with tempfile.TemporaryDirectory() as td:
        token = desktop_export.register_export_directory(td)
        csv_result = reports.save_csv(1, 'devices.csv', 'IP,Name\n10.0.0.1,Router\n', token)
        csv_path = Path(csv_result['path'])
        assert csv_path.parent == Path(td).resolve()
        assert csv_path.suffix == '.csv'
        assert csv_path.exists()
        try:
            reports.save_csv(1, 'again.csv', 'a,b\n1,2\n', token)
            assert False, 'destination token must be one-time'
        except ValueError as exc:
            assert 'EXPIRED' in str(exc)

        token2 = desktop_export.register_export_directory(td)
        xlsx_result = reports.save_xlsx_from_csv(1, 'devices.xlsx', 'IP,Name\n10.0.0.1,Router\n', token2)
        xlsx_path = Path(xlsx_result['path'])
        assert xlsx_path.parent == Path(td).resolve()
        assert xlsx_path.suffix == '.xlsx'
        wb = load_workbook(xlsx_path, read_only=True)
        ws = wb.active
        assert ws['A1'].value == 'IP'
        assert ws['B2'].value == 'Router'
        wb.close()


def test_backend_has_xlsx_endpoint_and_permission_guard():
    main = text('webapi/main.py')
    sec = text('webapi/security37.py')
    reports = text('webapi/reports37.py')
    assert "@app.post('/api/reports/xlsx')" in main
    assert 'save_xlsx_from_csv' in main
    assert "r'/api/reports/xlsx', ('Admin','Analyst','Operator')" in sec
    assert 'destination_token' in main
    assert 'def save_xlsx_from_csv' in reports


def test_export_picker_http_path_does_not_depend_on_js_bridge(monkeypatch, tmp_path):
    import desktop_export
    from webapi import desktop70

    desktop_export.set_native_export_picker(lambda: str(tmp_path))
    monkeypatch.setattr(desktop70, '_desktop_only', lambda: None)
    monkeypatch.setattr(desktop70.security37, 'require_role', lambda request, *roles: {'id': 1, 'role': 'Admin'})

    result = desktop70.choose_export_directory_http(object())
    assert result['cancelled'] is False
    assert result['token']
    assert desktop_export.consume_export_directory(result['token']) == tmp_path.resolve()
    desktop_export.set_native_export_picker(None)


def test_export_picker_cancel_is_normal_and_does_not_open_second_picker(monkeypatch):
    import desktop_export

    called = {'fallback': 0}
    desktop_export.set_native_export_picker(lambda: None)
    monkeypatch.setattr(desktop_export, '_powershell_folder_picker', lambda: called.__setitem__('fallback', called['fallback'] + 1))
    result = desktop_export.choose_export_directory()
    assert result == {'cancelled': True}
    assert called['fallback'] == 0
    desktop_export.set_native_export_picker(None)


def test_export_picker_falls_back_when_native_picker_raises(monkeypatch, tmp_path):
    import desktop_export

    def broken():
        raise RuntimeError('bridge unavailable')
    desktop_export.set_native_export_picker(broken)
    monkeypatch.setattr(desktop_export, '_powershell_folder_picker', lambda: str(tmp_path))
    result = desktop_export.choose_export_directory()
    assert result['cancelled'] is False
    assert desktop_export.consume_export_directory(result['token']) == tmp_path.resolve()
    desktop_export.set_native_export_picker(None)


def test_frontend_uses_desktop_http_picker_first():
    from pathlib import Path
    app = (Path(__file__).resolve().parents[1] / 'webapi/static/app.js').read_text(encoding='utf-8')
    chooser = app.split('async function chooseExportDirectory(){', 1)[1].split('function exportSavedText', 1)[0]
    assert "api('/desktop/export-directory',{method:'POST'})" in chooser
    assert chooser.index("api('/desktop/export-directory'") < chooser.index('window.pywebview')
    assert 'Hộp chọn thư mục Windows chưa sẵn sàng' not in chooser
